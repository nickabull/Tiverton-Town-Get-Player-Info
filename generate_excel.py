from __future__ import annotations

import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
DATA = BASE / "data.json"
OUT = BASE / "Tiverton_Town_SL1_Fixtures_2026-27.xlsx"

YELLOW = "FFF200"
DARK = "141414"
WHITE = "FFFFFF"
GRID = "8A8A8A"
RED = "D71920"
GREY = "808080"


def player_text(pd: dict, is_sub: bool) -> str:
    if not pd:
        return ""
    bits: list[str] = []
    if pd.get("captain"):
        bits.append("C")
    count = int(pd.get("goalCount") or 0)
    mins = [str(x) for x in (pd.get("goalMinutes") or []) if str(x)]
    if count:
        g = "G" if count == 1 else f"{count}G"
        if mins:
            g += " " + ", ".join(f"{m}'" for m in mins)
        bits.append(g)
    minute = pd.get("onMinute") if is_sub else pd.get("offMinute")
    if minute:
        bits.append(f"{minute}'")
    name = str(pd.get("name") or "")
    return name + (f" ({' / '.join(bits)})" if bits else "")


def main() -> None:
    payload = json.loads(DATA.read_text(encoding="utf-8-sig"))
    fixtures = sorted([x for x in payload.get("fixtures", []) if x.get("result")], key=lambda x: x.get("date", ""))

    wb = Workbook()
    ws = wb.active
    ws.title = "Fixtures"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D6"

    last_col = 23
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws["A1"] = "TIVERTON TOWN FC — 2026/27 SL1 FIXTURE MATRIX"
    ws["A1"].fill = PatternFill("solid", fgColor=DARK)
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color=YELLOW)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A2"] = "Starting XI in positional order • captain, goals and substitution minutes shown in brackets • unused subs in grey"
    ws["A2"].fill = PatternFill("solid", fgColor=DARK)
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color=WHITE)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    ws["A3"] = f"Source: {payload.get('source','')}  |  Updated: {payload.get('updated','')}"
    ws["A3"].font = Font(name="Arial", size=9, color="666666")
    ws["A3"].alignment = Alignment(horizontal="left")

    headers = ["Date", "Opponent", "Score"] + [str(i) for i in range(1, 12)] + ["SUBS"] + [f"S{i}" for i in range(1, 9)]
    for col, value in enumerate(headers, 1):
        c = ws.cell(5, col, value)
        c.fill = PatternFill("solid", fgColor=YELLOW)
        c.font = Font(name="Arial", size=10, bold=True, color="000000")
        c.alignment = Alignment(horizontal="center", vertical="center")

    thin = Side(style="thin", color=GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for idx, fx in enumerate(fixtures, start=6):
        fill_hex = YELLOW if (idx - 6) % 2 == 0 else WHITE
        row_fill = PatternFill("solid", fgColor=fill_hex)
        for col in range(1, last_col + 1):
            cell = ws.cell(idx, col)
            cell.fill = row_fill
            cell.border = border
            cell.font = Font(name="Arial", size=10, color="000000")
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        date = str(fx.get("date", ""))
        if len(date) == 10:
            date = f"{date[8:10]}/{date[5:7]}"
        ws.cell(idx, 1, date)
        opp = str(fx.get("opponent", ""))
        if str(fx.get("venue", "")).upper() == "H":
            opp = opp.upper()
        ws.cell(idx, 2, opp)
        ws.cell(idx, 3, str(fx.get("result", "")))
        for col in (1, 2, 3):
            ws.cell(idx, col).font = Font(name="Arial", size=10, bold=True)
        ws.cell(idx, 3).alignment = Alignment(horizontal="center", vertical="center")

        starters = fx.get("starterDetails") or []
        for i in range(11):
            if i < len(starters):
                ws.cell(idx, 4 + i, player_text(starters[i], False))

        ws.cell(idx, 15, "SUBS:")
        ws.cell(idx, 15).font = Font(name="Arial", size=10, bold=True)
        subs = fx.get("subDetails") or []
        for i in range(8):
            if i < len(subs):
                cell = ws.cell(idx, 16 + i, player_text(subs[i], True))
                if not bool(subs[i].get("used")):
                    cell.font = Font(name="Arial", size=10, color=GREY)
        ws.row_dimensions[idx].height = 36

    widths = {1: 9, 2: 24, 3: 9, 15: 8}
    for c in range(4, 15):
        widths[c] = 24
    for c in range(16, 24):
        widths[c] = 24
    for c, width in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = width

    # Keep dates and scores as literal text. Highlighting one side of a score within a
    # cell is not portable in openpyxl, so keep the score bold and centered in Excel.
    for row in range(6, 6 + len(fixtures)):
        ws.cell(row, 1).number_format = "@"
        ws.cell(row, 3).number_format = "@"

    wb.save(OUT)
    print(f"Wrote {OUT.name}")


if __name__ == "__main__":
    main()
